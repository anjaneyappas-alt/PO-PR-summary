import streamlit as st
import pandas as pd
import numpy as np
import io
import requests
from datetime import datetime, timedelta
from concurrent.futures import ThreadPoolExecutor, as_completed

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Page Config
st.set_page_config(page_title="PO vs PR Summary", page_icon="📊", layout="wide")

# Excel Grid Styling + High-Contrast KPI Cards
st.markdown("""
    <style>
        #MainMenu {visibility: hidden;}
        header {visibility: hidden;}
        footer {visibility: hidden;}
        [data-testid="stHeader"] {display: none;}
        
        .main {
            background-color: #ffffff !important;
        }

        [data-testid="stMetric"] {
            background-color: #f8f9fa !important;
            border: 1px solid #dee2e6 !important;
            padding: 12px 16px !important;
            border-radius: 8px !important;
            box-shadow: 0 1px 3px rgba(0,0,0,0.05) !important;
        }

        [data-testid="stMetricValue"] {
            font-size: 22px !important;
            font-weight: 700 !important;
            color: #111111 !important;
        }
        
        [data-testid="stMetricLabel"] {
            font-size: 13px !important;
            font-weight: 600 !important;
            color: #495057 !important;
        }

        .stTable {
            background-color: #ffffff !important;
            font-family: "Segoe UI", Arial, sans-serif !important;
            font-size: 13px !important;
            color: #000000 !important;
        }
        
        .stTable table {
            border-collapse: collapse !important;
            width: 100% !important;
            border: 1px solid #d9d9d9 !important;
        }
        
        .stTable th {
            background-color: #e6e6e6 !important;
            color: #000000 !important;
            font-weight: bold !important;
            text-align: center !important;
            border: 1px solid #bfbfbf !important;
            padding: 6px 10px !important;
            white-space: nowrap !important;
        }
        
        .stTable td {
            border: 1px solid #d9d9d9 !important;
            padding: 6px 10px !important;
            color: #000000 !important;
            white-space: nowrap !important;
        }
        
        .stTable tr:nth-child(even) {
            background-color: #f9f9f9 !important;
        }
    </style>
""", unsafe_allow_html=True)

st.title("📊 PO vs PR Excel View")

# Control Panel
control_col1, control_col2 = st.columns([1, 2])
with control_col1:
    data_source = st.radio("📡 Data Mode:", options=["Upload Files", "Live Zoho API Sync"], horizontal=True)
with control_col2:
    selected_window = st.radio(
        "⏱ Select PR Lookback Window:",
        options=["Last 15 Hours", "Last 24 Hours"],
        horizontal=True,
        index=0
    )

hours_window = 15 if selected_window == "Last 15 Hours" else 24


# --- ZOHO API HELPER FUNCTIONS ---
def get_zoho_access_token():
    client_id = st.secrets["zoho"]["client_id"]
    client_secret = st.secrets["zoho"]["client_secret"]
    refresh_token = st.secrets["zoho"]["refresh_token"]
    accounts_url = st.secrets["zoho"].get("accounts_url", "https://accounts.zoho.in")

    url = f"{accounts_url}/oauth/v2/token"
    payload = {
        "refresh_token": refresh_token,
        "client_id": client_id,
        "client_secret": client_secret,
        "grant_type": "refresh_token"
    }
    headers = {
        "Content-Type": "application/x-www-form-urlencoded"
    }
    response = requests.post(url, data=payload, headers=headers)
    res_data = response.json()
    if "access_token" in res_data:
        return res_data["access_token"]
    else:
        raise Exception(f"Failed to refresh Zoho Token: {res_data}")

def fetch_single_po_details(po_id, headers, domain, org_id):
    try:
        url = f"https://www.zohoapis.{domain}/inventory/v1/purchaseorders/{po_id}"
        res = requests.get(url, headers=headers, params={"organization_id": org_id}).json()
        if "purchaseorder" in res:
            po = res["purchaseorder"]
            line_items = po.get("line_items", [])
            total_qty = sum(float(item.get("quantity", 0)) for item in line_items)
            po["calculated_po_qty"] = total_qty
            return po
    except Exception:
        pass
    return None

def fetch_zoho_data_last_15_days():
    access_token = get_zoho_access_token()
    org_id = st.secrets["zoho"]["organization_id"]
    domain = st.secrets["zoho"].get("domain", "zoho.in")

    headers = {"Authorization": f"Zoho-oauthtoken {access_token}"}
    date_15_days_ago = (datetime.now() - timedelta(days=15)).strftime("%Y-%m-%d")
    date_90_days_ago = (datetime.now() - timedelta(days=90)).strftime("%Y-%m-%d")

    # 1. Fetch PRs from last 15 days first
    pr_url = f"https://www.zohoapis.{domain}/inventory/v1/purchasereceives"
    pr_params = {
        "organization_id": org_id, 
        "date_after": date_15_days_ago,
        "sort_column": "created_time",
        "sort_order": "D"
    }
    pr_res = requests.get(pr_url, headers=headers, params=pr_params).json()

    if "purchasereceives" not in pr_res:
        raise Exception(f"Zoho Purchase Receives API error: {pr_res}")

    pr_list = pr_res.get("purchasereceives", [])
    df_pr = pd.DataFrame(pr_list)

    # Find unique PO numbers referenced in these PRs
    po_ref_col = find_first_existing_col(df_pr, ['purchaseorder_number', 'PO Number', 'po_number', 'purchaseorder_no'])
    target_po_numbers = set(df_pr[po_ref_col].dropna().unique()) if po_ref_col in df_pr.columns else set()

    # 2. Fetch POs list (90-day window)
    po_url = f"https://www.zohoapis.{domain}/inventory/v1/purchaseorders"
    po_params = {
        "organization_id": org_id, 
        "date_after": date_90_days_ago,
        "sort_column": "date",
        "sort_order": "D"
    }
    po_res = requests.get(po_url, headers=headers, params=po_params).json()

    if "purchaseorders" not in po_res:
        raise Exception(f"Zoho Purchase Orders API error: {po_res}")

    po_summary_list = po_res.get("purchaseorders", [])

    # Filter POs that match our PRs
    po_to_fetch = [
        po for po in po_summary_list 
        if po.get("purchaseorder_number") in target_po_numbers or not target_po_numbers
    ]

    # Fast parallel fetching of line items for active POs
    detailed_pos = []
    with ThreadPoolExecutor(max_workers=10) as executor:
        futures = [
            executor.submit(fetch_single_po_details, po["purchaseorder_id"], headers, domain, org_id)
            for po in po_to_fetch if "purchaseorder_id" in po
        ]
        for future in as_completed(futures):
            result = future.result()
            if result:
                detailed_pos.append(result)

    df_po = pd.DataFrame(detailed_pos) if detailed_pos else pd.DataFrame(po_summary_list)

    return df_po, df_pr


# Data Input Handling
df_po_raw, df_pr_raw = None, None

if data_source == "Upload Files":
    col1, col2 = st.columns(2)
    with col1:
        po_file = st.file_uploader("Upload PO Monthly Data (Excel)", type=['xlsx', 'xls'])
    with col2:
        pr_file = st.file_uploader("Upload PR Monthly Data (Excel)", type=['xlsx', 'xls'])
    if po_file and pr_file:
        df_po_raw = pd.read_excel(po_file)
        df_pr_raw = pd.read_excel(pr_file)

else:
    st.markdown("<br>", unsafe_allow_html=True)
    if st.button("🔄 Fetch Last 15 Days Data from Zoho", type="primary"):
        try:
            with st.spinner("Fetching live PO & PR line-item data from Zoho..."):
                df_po_raw, df_pr_raw = fetch_zoho_data_last_15_days()
                st.session_state["raw_po"] = df_po_raw
                st.session_state["raw_pr"] = df_pr_raw
                st.success("Successfully fetched live data from Zoho!")
        except Exception as e:
            st.error(f"Zoho API Error: {e}")

    if "raw_po" in st.session_state and "raw_pr" in st.session_state:
        df_po_raw = st.session_state["raw_po"]
        df_pr_raw = st.session_state["raw_pr"]

# Helper for robust column lookup
def find_first_existing_col(df, candidates, default=None):
    for col in candidates:
        if col in df.columns:
            return col
    return default

# Process Data
if df_po_raw is not None and df_pr_raw is not None:
    try:
        with st.spinner("Processing reconciliation..."):
            df_po = df_po_raw.copy()
            df_pr = df_pr_raw.copy()

            df_po.columns = df_po.columns.str.strip()
            df_pr.columns = df_pr.columns.str.strip()

            po_num_col = find_first_existing_col(df_po, ['purchaseorder_number', 'Purchase Order Number', 'PO Number', 'purchaseorder_no'])
            po_qty_col = find_first_existing_col(df_po, ['calculated_po_qty', 'QuantityOrdered', 'Quantity', 'quantity', 'quantity_ordered'])

            pr_no_col = find_first_existing_col(df_pr, ['purchasereceive_number', 'receive_number', 'Receive Number', 'PR Number'])
            qty_pr_col = find_first_existing_col(df_pr, ['quantity', 'quantity_received', 'Quantity Received', 'total_quantity', 'Quantity'])
            
            po_ref_in_pr = find_first_existing_col(df_pr, ['purchaseorder_number', 'PO Number', 'po_number', 'purchaseorder_no'])

            time_col = find_first_existing_col(df_pr, ['created_time', 'CreatedTime', 'date', 'Receive Date'])

            df_pr_clean = df_pr.dropna(subset=[po_ref_in_pr]).copy() if po_ref_in_pr else df_pr.copy()
            
            if qty_pr_col in df_pr_clean.columns:
                df_pr_clean['Clean_PR_Qty'] = pd.to_numeric(df_pr_clean[qty_pr_col], errors='coerce').fillna(0)
            else:
                df_pr_clean['Clean_PR_Qty'] = 0

            df_pr_clean['DT'] = pd.to_datetime(df_pr_clean[time_col], errors='coerce') if time_col else pd.NaT
            
            latest_pr_time = df_pr_clean['DT'].max() if 'DT' in df_pr_clean.columns else pd.NaT
            
            if pd.notna(latest_pr_time):
                cutoff_time = latest_pr_time - pd.Timedelta(hours=hours_window)
                df_today = df_pr_clean[df_pr_clean['DT'] >= cutoff_time]
                df_prior = df_pr_clean[df_pr_clean['DT'] < cutoff_time]
            else:
                df_today = df_pr_clean
                df_prior = pd.DataFrame(columns=df_pr_clean.columns)

            vendor_col_pr = find_first_existing_col(df_pr_clean, ['vendor_name', 'Vendor Name'])
            vendor_col_po = find_first_existing_col(df_po, ['vendor_name', 'Vendor Name'])

            today_summary = df_today.groupby(po_ref_in_pr).agg(
                PR_no=(pr_no_col, lambda x: " & ".join(sorted(x.dropna().astype(str).unique()))),
                PR_Qty=('Clean_PR_Qty', 'sum'),
                PR_Date=('DT', 'max'),
                Vendor_PR=(vendor_col_pr, 'first') if vendor_col_pr in df_today.columns else (pr_no_col, 'first')
            ).reset_index()

            prior_summary = df_prior.groupby(po_ref_in_pr).agg(
                PRMTD=('Clean_PR_Qty', 'sum')
            ).reset_index()

            df_po_clean = df_po.dropna(subset=[po_num_col]).copy() if po_num_col else df_po.copy()
            
            if po_qty_col and po_qty_col in df_po_clean.columns:
                df_po_clean['Clean_PO_Qty'] = pd.to_numeric(df_po_clean[po_qty_col], errors='coerce').fillna(0)
            else:
                df_po_clean['Clean_PO_Qty'] = 0

            po_summary = df_po_clean.groupby(po_num_col).agg(
                PO_Qty=('Clean_PO_Qty', 'sum'),
                Vendor_PO=(vendor_col_po, 'first') if vendor_col_po in df_po_clean.columns else (po_num_col, 'first')
            ).reset_index()

            merged = pd.merge(today_summary, po_summary, left_on=po_ref_in_pr, right_on=po_num_col, how='left')
            merged = pd.merge(merged, prior_summary, on=po_ref_in_pr, how='left')

            if 'PR_Date' in merged.columns and pd.api.types.is_datetime64_any_dtype(merged['PR_Date']):
                merged = merged.sort_values(by='PR_Date', ascending=True).reset_index(drop=True)
                merged['Date'] = merged['PR_Date'].dt.strftime('%d-%m-%y').fillna('')
            else:
                merged['Date'] = ''

            merged['Vendor Name'] = merged['Vendor_PR'].fillna(merged['Vendor_PO']).fillna('')
            
            merged['PO Qty'] = pd.to_numeric(merged['PO_Qty'], errors='coerce').fillna(0).astype(int)
            merged['PR Qty'] = pd.to_numeric(merged['PR_Qty'], errors='coerce').fillna(0).astype(int)
            merged['PRMTD'] = pd.to_numeric(merged['PRMTD'], errors='coerce').fillna(0).astype(int)
            
            merged['Total Received'] = merged['PR Qty'] + merged['PRMTD']
            merged['Pending Qty'] = (merged['PO Qty'] - merged['Total Received']).apply(lambda x: int(x) if x > 0 else 0)
            merged['Sl.no'] = range(1, len(merged) + 1)

            total_pos_cnt = len(merged)
            total_po = int(merged['PO Qty'].sum())
            total_pr = int(merged['PR Qty'].sum())
            total_prmtd = int(merged['PRMTD'].sum())
            total_pending = int(merged['Pending Qty'].sum())
            total_fr_val = ((total_pr + total_prmtd) / total_po * 100) if total_po > 0 else 0

            expected_headers = ["Sl.no", "Date", po_ref_in_pr, "PR_no", "Vendor Name", "PO Qty", "PR Qty", "PRMTD", "Pending Qty"]
            final_df = merged[expected_headers].rename(columns={po_ref_in_pr: 'PO No', 'PR_no': 'PR No', 'PR Qty': 'Today PR Qty'})
            
            fr_numeric = np.where(final_df['PO Qty'] > 0, ((final_df['Today PR Qty'] + final_df['PRMTD']) / final_df['PO Qty']) * 100, 0)
            final_df['PO FR %'] = np.round(fr_numeric).astype(int).astype(str) + '%'

            total_row = pd.DataFrame([{
                "Sl.no": "",
                "Date": "",
                "PO No": "",
                "PR No": "",
                "Vendor Name": "Total",
                "PO Qty": total_po,
                "Today PR Qty": total_pr,
                "PRMTD": total_prmtd,
                "Pending Qty": total_pending,
                "PO FR %": f"{int(round(total_fr_val))}%"
            }])

            display_df = pd.concat([final_df, total_row], ignore_index=True)

            st.session_state["processed_df"] = display_df
            st.session_state["kpi_metrics"] = (total_pos_cnt, total_po, total_pr, total_prmtd, total_pending, total_fr_val)
            st.session_state["anchor_time"] = latest_pr_time
            st.session_state["window_hours"] = hours_window

    except Exception as e:
        st.error(f"Error processing data: {e}")

# Render UI
if "processed_df" in st.session_state:
    display_df = st.session_state["processed_df"]
    total_pos_cnt, total_po, total_pr, total_prmtd, total_pending, total_fr_val = st.session_state["kpi_metrics"]
    anchor_time = st.session_state.get("anchor_time", None)
    active_window = st.session_state.get("window_hours", 15)

    st.markdown("### 🎯 Total Summary")
    if anchor_time and pd.notna(anchor_time):
        st.markdown(f"🕒 **Last PR Created Time:** `{anchor_time.strftime('%d-%b-%Y %H:%M:%S')}` | **{active_window}-Hour Window:** `{ (anchor_time - pd.Timedelta(hours=active_window)).strftime('%d-%b-%Y %H:%M:%S') }` onwards")

    kpi1, kpi2, kpi3, kpi4, kpi5, kpi6 = st.columns(6)
    kpi1.metric("Total POs", f"{total_pos_cnt:,}")
    kpi2.metric("PO Qty", f"{total_po:,}")
    kpi3.metric(f"Today PR Qty ({active_window}h)", f"{total_pr:,}")
    kpi4.metric("PRMTD Qty", f"{total_prmtd:,}")
    kpi5.metric("Pending Qty", f"{total_pending:,}")
    kpi6.metric("Fill Rate", f"{total_fr_val:.1f}%")

    st.markdown("---")

    def highlight_excel_cells(df):
        styles = pd.DataFrame('', index=df.index, columns=df.columns)
        for idx, row in df.iterrows():
            if str(row['Vendor Name']) == 'Total':
                styles.loc[idx, :] = 'background-color: #f4b084; font-weight: bold; color: black; border-top: 2px solid black; border-bottom: 2px double black'
            else:
                try:
                    if float(row['Pending Qty']) > 0:
                        styles.loc[idx, 'Pending Qty'] = 'background-color: #fff3cd; color: #856404; font-weight: bold;'
                except (ValueError, TypeError):
                    pass
        return styles

    styled_df = display_df.style.apply(highlight_excel_cells, axis=None)
    st.table(styled_df)

    def generate_excel_file(df):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "PO vs PR Summary"
        ws.views.sheetView[0].showGridLines = True

        header_fill = PatternFill(start_color="E6E6E6", end_color="E6E6E6", fill_type="solid")
        zebra_fill = PatternFill(start_color="F9F9F9", end_color="F9F9F9", fill_type="solid")
        white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        yellow_pending_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
        orange_total_fill = PatternFill(start_color="F4B084", end_color="F4B084", fill_type="solid")

        thin_side = Side(border_style="thin", color="D9D9D9")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        
        total_top_side = Side(border_style="thin", color="000000")
        total_bottom_side = Side(border_style="double", color="000000")
        total_border = Border(left=thin_side, right=thin_side, top=total_top_side, bottom=total_bottom_side)

        font_header = Font(name="Segoe UI", size=11, bold=True, color="000000")
        font_regular = Font(name="Segoe UI", size=11, bold=False, color="000000")
        font_pending = Font(name="Segoe UI", size=11, bold=True, color="856404")
        font_total = Font(name="Segoe UI", size=11, bold=True, color="000000")

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")

        headers = list(df.columns)
        ws.append(headers)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = font_header
            cell.alignment = align_center
            cell.border = Border(left=Side(border_style="thin", color="BFBFBF"),
                                 right=Side(border_style="thin", color="BFBFBF"),
                                 top=Side(border_style="thin", color="BFBFBF"),
                                 bottom=Side(border_style="thin", color="BFBFBF"))

        pending_col_idx = headers.index('Pending Qty') + 1 if 'Pending Qty' in headers else None

        for r_idx, row_data in enumerate(df.values, start=2):
            ws.append(list(row_data))
            is_total_row = (str(row_data[headers.index('Vendor Name')]) == 'Total') if 'Vendor Name' in headers else False

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=r_idx, column=col_idx)
                val = cell.value

                if is_total_row:
                    cell.fill = orange_total_fill
                    cell.font = font_total
                    cell.border = total_border
                else:
                    cell.fill = zebra_fill if (r_idx % 2 == 0) else white_fill
                    cell.font = font_regular
                    cell.border = thin_border

                    if col_idx == pending_col_idx:
                        try:
                            if float(val) > 0:
                                cell.fill = yellow_pending_fill
                                cell.font = font_pending
                        except (ValueError, TypeError):
                            pass

                if headers[col_idx - 1] in ['Sl.no', 'Date', 'PO No', 'PR No', 'PO FR %']:
                    cell.alignment = align_center
                elif headers[col_idx - 1] in ['Vendor Name']:
                    cell.alignment = align_left
                else:
                    cell.alignment = align_right

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    st.markdown("---")
    excel_bytes = generate_excel_file(display_df)
    st.download_button(
        label="📥 Download Summary as Excel (.xlsx)",
        data=excel_bytes,
        file_name=f"PO_PR_Summary_{active_window}h.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
